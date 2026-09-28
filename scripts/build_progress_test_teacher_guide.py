from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUTPUT = Path("/Users/nicholasbeeson/aptis-grammar-tool/docs/teacher-guides/Teacher_Guide_to_Progress_Tests.docx")
SCREENSHOT_DIR = OUTPUT.parent / "progress-test-screenshots"
NAVY = "17315E"
PALE_BLUE = "EAF1FB"
MID_BLUE = "315A91"
LIGHT_GREY = "F3F5F7"
BORDER = "D9D9D9"
BLACK = RGBColor(0, 0, 0)
SOFT = RGBColor(75, 82, 92)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_borders(cell, color=BORDER, size="6"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:color"), color)


def set_cell_margins(cell, top=120, start=140, bottom=120, end=140):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn("w:" + margin))
        if node is None:
            node = OxmlElement("w:" + margin)
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def add_page_number(paragraph):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instruction, end])


def configure_styles(doc):
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(11)
    normal.font.color.rgb = BLACK
    normal.paragraph_format.space_after = Pt(7)
    normal.paragraph_format.line_spacing = 1.12

    title = styles["Title"]
    title.font.name = "Aptos Display"
    title.font.size = Pt(30)
    title.font.bold = True
    title.font.color.rgb = BLACK
    title.paragraph_format.space_after = Pt(10)

    for style_name, size, before, after in (
        ("Heading 1", 21, 14, 8),
        ("Heading 2", 15, 11, 5),
        ("Heading 3", 12, 8, 4),
    ):
        style = styles[style_name]
        style.font.name = "Aptos Display"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = BLACK
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    if "Screenshot Caption" not in styles:
        caption = styles.add_style("Screenshot Caption", WD_STYLE_TYPE.PARAGRAPH)
    else:
        caption = styles["Screenshot Caption"]
    caption.font.name = "Aptos"
    caption.font.size = Pt(9)
    caption.font.italic = True
    caption.font.color.rgb = SOFT
    caption.paragraph_format.space_before = Pt(4)
    caption.paragraph_format.space_after = Pt(8)
    caption.paragraph_format.keep_with_next = True

    if "Screenshot Placeholder" not in styles:
        placeholder = styles.add_style("Screenshot Placeholder", WD_STYLE_TYPE.PARAGRAPH)
    else:
        placeholder = styles["Screenshot Placeholder"]
    placeholder.font.name = "Aptos"
    placeholder.font.size = Pt(9)
    placeholder.font.color.rgb = RGBColor(255, 255, 255)
    placeholder.paragraph_format.space_after = Pt(0)
    placeholder.paragraph_format.keep_together = True


def add_header_footer(doc):
    for section in doc.sections:
        header = section.header.paragraphs[0]
        header.text = "Teacher guide to progress tests"
        header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        header_run = header.runs[0]
        header_run.font.name = "Aptos"
        header_run.font.size = Pt(8.5)
        header_run.font.color.rgb = SOFT

        footer = section.footer.paragraphs[0]
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = footer.add_run("Seif English Academy   •   ")
        run.font.name = "Aptos"
        run.font.size = Pt(8)
        run.font.color.rgb = SOFT
        add_page_number(footer)


def add_step(doc, number, title, body):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(5)
    paragraph.paragraph_format.space_after = Pt(5)
    paragraph.paragraph_format.keep_together = True
    number_run = paragraph.add_run(f"{number}  ")
    number_run.bold = True
    number_run.font.color.rgb = RGBColor(49, 90, 145)
    title_run = paragraph.add_run(title)
    title_run.bold = True
    paragraph.add_run(f"  {body}")


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(4)
        p.add_run(item)


def add_screenshot(doc, number, caption, width=6.65):
    cap = doc.add_paragraph(caption, style="Screenshot Caption")
    cap.paragraph_format.keep_with_next = True
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_together = True
    p.paragraph_format.keep_with_next = False
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(8)
    screenshot_files = {
        1: "final-1.jpg",
        2: "final-2.jpg",
        3: "final-3.jpg",
        4: "final-4.jpg",
        5: "final-5.jpg",
        6: "final-6.jpg",
        7: "final-7.jpg",
    }
    screenshot_path = SCREENSHOT_DIR / screenshot_files.get(number, f"screenshot-{number}.jpg")
    if not screenshot_path.exists():
        raise FileNotFoundError(f"Missing guide screenshot: {screenshot_path}")
    p.add_run().add_picture(str(screenshot_path), width=Inches(width))


def add_button_table(doc):
    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.columns[0].width = Inches(1.8)
    table.columns[1].width = Inches(4.8)
    headers = ("Button", "Use")
    for i, text in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.width = table.columns[i].width
        set_cell_shading(cell, NAVY)
        set_cell_borders(cell)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(text)
        r.bold = True
        r.font.color.rgb = RGBColor(255, 255, 255)

    rows = [
        ("Open main paper", "Lets students begin the timed grammar, vocabulary, pronunciation, and reading paper."),
        ("Open listening", "Makes the listening stage available. Open it only when you are ready for the class to start together."),
        ("View exam", "Opens the fixed test content so you can check the paper before the session."),
        ("Do session", "Runs a teacher preview. No student submission is recorded."),
        ("Live timings", "Shows how many students have not started, are in the main paper, are waiting, are listening, or have finished."),
        ("Review submissions", "Opens completed attempts for marking and final review."),
    ]
    for row_index, (button, use) in enumerate(rows, start=1):
        cells = table.add_row().cells
        fill = "FFFFFF" if row_index % 2 else PALE_BLUE
        for i, text in enumerate((button, use)):
            cells[i].width = table.columns[i].width
            set_cell_shading(cells[i], fill)
            set_cell_borders(cells[i])
            set_cell_margins(cells[i], top=105, bottom=105)
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cells[i].paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(text)
            if i == 0:
                r.bold = True
                r.font.color.rgb = RGBColor(36, 70, 116)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def build_document():
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.7)
    section.right_margin = Inches(0.7)

    configure_styles(doc)
    add_header_footer(doc)

    title = doc.add_paragraph("Teacher Guide to Progress Tests", style="Title")
    title.paragraph_format.space_before = Pt(34)
    subtitle = doc.add_paragraph("Set up run and review a teacher controlled course test")
    subtitle.style = doc.styles["Subtitle"]
    subtitle.runs[0].font.color.rgb = BLACK
    subtitle.runs[0].font.size = Pt(16)

    intro = doc.add_paragraph()
    intro.paragraph_format.space_before = Pt(18)
    intro.paragraph_format.space_after = Pt(14)
    intro.add_run(
        "This guide shows teachers how to create a progress test session, assign students, open each stage, monitor the class, run the listening paper, and review submissions. It uses teacher controlled mode, which is the best choice when the class completes the test together."
    )

    doc.add_heading("The complete workflow", level=1)
    add_step(doc, "1", "Prepare", "Choose the fixed test template, session title, class, students, dates, and optional PIN.")
    add_step(doc, "2", "Open the main paper", "Start the timed written paper when everyone is ready.")
    add_step(doc, "3", "Open listening", "Release listening after the written paper has been submitted and the class is ready.")
    add_step(doc, "4", "Review", "Open completed attempts, check marks that need teacher judgement, and confirm the result.")

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    r = p.add_run("Before class")
    r.bold = True
    p.add_run("  Preview the paper, confirm the assigned students, test the room audio, and keep the Teacher Tools page open.")

    doc.add_page_break()
    doc.add_heading("Open the course test setup", level=1)
    doc.add_paragraph(
        "Sign in with a teacher or admin account. Open Teacher Tools, then select Set up course tests. The panel contains the session form on the left and your existing sessions on the right."
    )
    add_screenshot(doc, 1, "Figure 1  Open Set up course tests from Teacher Tools")

    doc.add_heading("Choose the test and title", level=2)
    add_bullets(doc, [
        "Choose the progress test that matches the course and units your class has completed.",
        "Give the session a title that students will recognise, such as B1 Progress Test October.",
        "Admins should check the Owner teacher field. Teachers normally do not see this field.",
    ])
    add_screenshot(doc, 2, "Figure 2  Choose the fixed test template and enter a clear session title", width=5.8)

    doc.add_page_break()
    doc.add_heading("Choose access and session control", level=1)
    doc.add_paragraph(
        "Use the class filter to narrow the student list. A session PIN is optional. Generate one when students need an extra check before entering the test."
    )
    add_screenshot(doc, 3, "Figure 3  Set the class access mode and session window", width=5.8)

    doc.add_heading("Teacher controlled and self controlled modes", level=2)
    table = doc.add_table(rows=1, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    widths = (Inches(1.45), Inches(2.55), Inches(2.55))
    for i, (text, width) in enumerate(zip(("Mode", "Use it when", "What happens"), widths)):
        cell = table.rows[0].cells[i]
        cell.width = width
        set_cell_shading(cell, NAVY)
        set_cell_borders(cell)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(text)
        r.bold = True
        r.font.color.rgb = RGBColor(255, 255, 255)
    mode_rows = [
        ("Teacher controlled", "The class sits the test together.", "Students wait until you open the main paper and listening."),
        ("Self controlled", "Students complete the test independently during a wider window.", "The main paper and listening are already open during the session window."),
    ]
    for row_index, values in enumerate(mode_rows, start=1):
        cells = table.add_row().cells
        for i, value in enumerate(values):
            cells[i].width = widths[i]
            set_cell_shading(cells[i], "FFFFFF" if row_index % 2 else PALE_BLUE)
            set_cell_borders(cells[i])
            set_cell_margins(cells[i])
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cells[i].paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(value)
            if i == 0:
                r.bold = True

    doc.add_heading("Assign students and create the session", level=2)
    add_bullets(doc, [
        "Select a class filter, then use Select filtered or choose students individually.",
        "Set the available dates if the session should open and close automatically.",
        "Add brief teacher notes only when students need extra instructions.",
        "Check the selected student count, then select Create session.",
    ])

    doc.add_page_break()
    doc.add_heading("Prepare the session before students begin", level=1)
    doc.add_paragraph(
        "Find the session under My course test sessions. Teacher controlled sessions show separate controls for the main paper and listening. Both stages start locked."
    )
    add_screenshot(doc, 4, "Figure 4  Teacher controlled session buttons", width=5.6)

    add_button_table(doc)

    doc.add_heading("Preview the paper", level=2)
    doc.add_paragraph(
        "Select Do session to run the test exactly as a student would. The page clearly identifies teacher preview mode, and no student attempt or submission is recorded."
    )
    add_screenshot(doc, 5, "Figure 5  Teacher preview is separate from student attempts", width=5.7)

    doc.add_page_break()
    doc.add_heading("Run the main paper", level=1)
    add_step(doc, "1", "Gather the class", "Ask students to sign in and open Your Class. Confirm that the correct progress test appears.")
    add_step(doc, "2", "Open the main paper", "Select Open main paper on the session card when the room is ready.")
    add_step(doc, "3", "Students begin", "Students open the session and start the timer. They can move between sections and their answers save automatically.")
    add_step(doc, "4", "Monitor progress", "Use Live timings to see how many students have not started, are working, are waiting for listening, or have finished.")

    add_screenshot(doc, 6, "Figure 6  The teacher preview shows the same section layout students use", width=5.7)

    doc.add_heading("When the main paper ends", level=2)
    doc.add_paragraph(
        "Students submit the main paper and move to a waiting screen. Keep listening locked until the whole class is ready. A student who submits the main paper cannot return to grammar, vocabulary, pronunciation, or reading."
    )

    doc.add_page_break()
    doc.add_heading("Run the listening paper", level=1)
    add_step(doc, "1", "Prepare the room", "Check that students have headphones or that the room speakers are ready. Ask everyone to remain on the listening page.")
    add_step(doc, "2", "Open listening", "Select Open listening on the session card. Students can now start the listening stage.")
    add_step(doc, "3", "Allow both plays", "Each exercise gives reading time, plays the recording once, shows a clear countdown, and then plays it a second time automatically.")
    add_step(doc, "4", "Watch live timings", "The listening status shows each student’s current phase, including the first play, the pause before replay, the second play, and completion.")

    doc.add_heading("What students should know", level=2)
    add_bullets(doc, [
        "The second listen starts automatically. Students should not select Skip exercise early during the pause.",
        "If a browser blocks automatic audio, the student must select Play recording 1 or Play recording 2.",
        "Leaving early shows a warning that explains whether the first listen, second listen, or exercise is still incomplete.",
        "Submitting with unanswered listening questions shows a final warning.",
    ])

    add_screenshot(doc, 7, "Figure 7  Live timings shows the class position without opening individual work")

    doc.add_page_break()
    doc.add_heading("Review and finish the test", level=1)
    doc.add_paragraph(
        "After students submit listening, select Review submissions. Automatically marked items already have a score. Review any items that need teacher judgement, then save the final result."
    )
    add_bullets(doc, [
        "Use Live timings first if a student appears to be missing. They may still be in the main paper, waiting, or listening.",
        "Open Review submissions only after the student has submitted the full test.",
        "Check the section breakdown before confirming the final result.",
        "Keep the session for your records. Delete it only when you are certain it is no longer needed.",
    ])

    doc.add_heading("Classroom checklist", level=1)
    checklist = [
        "The correct template and unit range are selected",
        "The correct students are assigned",
        "The session is teacher controlled",
        "The main paper remains locked until the class is ready",
        "Listening remains locked until main paper submissions are complete",
        "Students know that every listening exercise plays twice",
        "The room audio has been tested",
        "Live timings is open during the test",
    ]
    for item in checklist:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(5)
        r = p.add_run("☐  ")
        r.font.size = Pt(13)
        p.add_run(item)

    doc.add_heading("If something goes wrong", level=2)
    troubleshooting = [
        ("A student cannot see the test", "Check that the student is assigned, the session window is open, and the correct stage has been released."),
        ("Audio does not start", "Ask the student to use the visible Play recording button. This appears when the browser blocks automatic playback."),
        ("A page was refreshed during listening", "The listening exercise pauses to protect the attempt. Check Live timings and decide with the student before continuing."),
        ("A student skipped early", "The session records an early listening finish state. Ask the student what they heard and review unanswered items before final marking."),
    ]
    for problem, action in troubleshooting:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(5)
        r = p.add_run(problem + "  ")
        r.bold = True
        p.add_run(action)

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build_document()
